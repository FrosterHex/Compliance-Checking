from pptx import Presentation
from pptx.util import Inches, Pt

def make_pptx(out_path="review2/Review2_Presentation.pptx"):
    prs = Presentation()

    # Title slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Project Title"
    subtitle.text = "Student Names — Guide Name"

    # Guide approval slide (Slide 2)
    content_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(content_layout)
    slide.shapes.title.text = "Guide Approval (REQUIRED)"
    tf = slide.shapes.placeholders[1].text_frame
    tf.text = "Paste guide approval text or screenshot here.\nName:\nEmail:\nDate:\n"

    # Other slides (from ppt_outline)
    titles = [
        "Problem Statement",
        "Motivation & Objectives",
        "Scope & Constraints",
        "System Architecture",
        "Technical Approach / Methodology",
        "Implementation Status",
        "Demo Plan / How to reproduce the demo",
        "Results / Preliminary Findings",
        "Gantt / Timeline / Next Steps",
        "Risks & Mitigations",
        "Conclusions",
        "Q&A",
    ]

    for t in titles:
        slide = prs.slides.add_slide(content_layout)
        slide.shapes.title.text = t
        tf = slide.shapes.placeholders[1].text_frame
        tf.text = "Notes:"

    prs.save(out_path)
    print(f"Wrote {out_path}")

if __name__ == "__main__":
    make_pptx()
